from pathlib import Path
import io
import re

import pandas as pd
import fitz
import pdfplumber

from docx import Document


# =========================================================
# SUPPORTED FILE TYPES
# =========================================================

SUPPORTED_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".pdf",
    ".docx",
    ".txt",
    ".json",
    ".parquet",
}


# =========================================================
# FILE EXTENSION
# =========================================================

def get_extension(filename: str) -> str:

    return Path(
        filename
    ).suffix.lower()


# =========================================================
# BASIC CLEANING
# =========================================================

def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:

    if df is None:
        return pd.DataFrame()

    df = df.copy()

    # Remove completely empty rows
    df = df.dropna(
        how="all"
    )

    # Remove completely empty columns
    df = df.dropna(
        axis=1,
        how="all"
    )

    if df.empty:
        return df

    # Clean column names
    cleaned_columns = []

    for column in df.columns:

        column_name = str(
            column
        ).strip()

        column_name = re.sub(
            r"\s+",
            "_",
            column_name
        )

        column_name = re.sub(
            r"_+",
            "_",
            column_name
        )

        cleaned_columns.append(
            column_name
        )

    df.columns = cleaned_columns

    # Clean string cells
    for column in df.columns:

        if df[column].dtype == "object":

            df[column] = (
                df[column]
                .astype(str)
                .str.strip()
                .replace(
                    {
                        "": pd.NA,
                        "nan": pd.NA,
                        "None": pd.NA,
                    }
                )
            )

    return df.reset_index(
        drop=True
    )


# =========================================================
# TYPE CONVERSION
# =========================================================

def convert_dataframe_types(
    df: pd.DataFrame
) -> pd.DataFrame:

    if df.empty:
        return df

    df = df.copy()

    for column in df.columns:

        series = df[column]

        # Do not convert obvious identifier columns
        normalized = (
            str(column)
            .strip()
            .lower()
        )

        id_column = (
            normalized == "id"
            or normalized.endswith("_id")
            or normalized.endswith("id")
            or normalized in {
                "employee_id",
                "customer_id",
                "transaction_id",
                "product_id",
                "order_id",
                "supplier_id",
                "patient_id",
                "machine_id",
            }
        )

        if id_column:
            continue

        # Numeric conversion
        if series.dtype == "object":

            numeric_candidate = (
                series.astype(str)
                .str.replace(
                    ",",
                    "",
                    regex=False
                )
                .str.replace(
                    "₹",
                    "",
                    regex=False
                )
                .str.replace(
                    "$",
                    "",
                    regex=False
                )
                .str.replace(
                    "%",
                    "",
                    regex=False
                )
                .str.strip()
            )

            numeric_values = pd.to_numeric(
                numeric_candidate,
                errors="coerce"
            )

            valid_count = numeric_values.notna().sum()
            total_count = series.notna().sum()

            if (
                total_count > 0
                and valid_count / total_count >= 0.80
            ):

                df[column] = numeric_values

                continue

        # Date conversion
        if series.dtype == "object":

            date_values = pd.to_datetime(
                series,
                errors="coerce"
            )

            valid_count = date_values.notna().sum()
            total_count = series.notna().sum()

            if (
                total_count > 0
                and valid_count / total_count >= 0.80
            ):

                df[column] = date_values

    return df


# =========================================================
# PDF HEADER CLEANING
# =========================================================

def clean_pdf_header(
    header
):

    if header is None:
        return None

    header = str(
        header
    ).strip()

    header = re.sub(
        r"\s+",
        "_",
        header
    )

    header = re.sub(
        r"_+",
        "_",
        header
    )

    if not header:
        return None

    return header


# =========================================================
# PDF TABLE CLEANING
# =========================================================

def clean_pdf_table(
    table
):

    if not table:
        return pd.DataFrame()

    cleaned_rows = []

    for row in table:

        if row is None:
            continue

        cleaned_row = []

        for cell in row:

            if cell is None:
                cleaned_row.append("")

            else:

                value = str(
                    cell
                ).strip()

                value = re.sub(
                    r"\s+",
                    " ",
                    value
                )

                cleaned_row.append(
                    value
                )

        # Skip completely empty rows
        if any(
            value != ""
            for value in cleaned_row
        ):

            cleaned_rows.append(
                cleaned_row
            )

    if len(cleaned_rows) < 2:
        return pd.DataFrame()

    # Determine maximum column count
    column_count = max(
        len(row)
        for row in cleaned_rows
    )

    normalized_rows = []

    for row in cleaned_rows:

        if len(row) < column_count:

            row = row + (
                [""] *
                (
                    column_count -
                    len(row)
                )
            )

        elif len(row) > column_count:

            row = row[:column_count]

        normalized_rows.append(
            row
        )

    header = normalized_rows[0]

    # Clean headers
    headers = []

    for index, value in enumerate(header):

        cleaned = clean_pdf_header(
            value
        )

        if not cleaned:

            cleaned = (
                f"Column_{index + 1}"
            )

        headers.append(
            cleaned
        )

    # Ensure unique column names
    unique_headers = []
    header_counts = {}

    for header in headers:

        if header not in header_counts:

            header_counts[header] = 0
            unique_headers.append(
                header
            )

        else:

            header_counts[header] += 1

            unique_headers.append(
                f"{header}_{header_counts[header]}"
            )

    data_rows = normalized_rows[1:]

    df = pd.DataFrame(
        data_rows,
        columns=unique_headers
    )

    return clean_dataframe(
        df
    )


# =========================================================
# PDF TABLE EXTRACTION
# =========================================================

def extract_pdf_tables(
    raw: bytes
):

    tables = []

    try:

        with pdfplumber.open(
            io.BytesIO(raw)
        ) as pdf:

            for page_number, page in enumerate(
                pdf.pages,
                start=1
            ):

                try:

                    page_tables = (
                        page.extract_tables()
                    )

                except Exception:

                    page_tables = []

                if not page_tables:
                    continue

                for table_index, table in enumerate(
                    page_tables,
                    start=1
                ):

                    df = clean_pdf_table(
                        table
                    )

                    if df.empty:
                        continue

                    df = convert_dataframe_types(
                        df
                    )

                    tables.append(
                        (
                            page_number,
                            table_index,
                            df
                        )
                    )

    except Exception:
        return []

    return tables


# =========================================================
# PDF TEXT EXTRACTION
# =========================================================

def extract_pdf_text(
    raw: bytes
):

    try:

        document = fitz.open(
            stream=raw,
            filetype="pdf"
        )

        text_parts = []

        for page in document:

            text = page.get_text()

            if text:

                text_parts.append(
                    text.strip()
                )

        document.close()

        return "\n".join(
            text_parts
        )

    except Exception:

        return ""


# =========================================================
# PDF TABLE QUALITY CHECK
# =========================================================

def is_useful_table(
    df: pd.DataFrame
) -> bool:

    if df is None:
        return False

    if df.empty:
        return False

    if len(df.columns) < 2:
        return False

    if len(df) < 1:
        return False

    # Avoid treating a simple one-column text block
    # as a business table.
    non_empty_cells = (
        df.notna()
        .sum()
        .sum()
    )

    return non_empty_cells >= 4


# =========================================================
# CSV LOADER
# =========================================================

def load_csv(
    raw: bytes
):

    df = pd.read_csv(
        io.BytesIO(raw)
    )

    df = clean_dataframe(
        df
    )

    df = convert_dataframe_types(
        df
    )

    return {
        "data": df
    }


# =========================================================
# EXCEL LOADER
# =========================================================

def load_excel(
    raw: bytes
):

    sheets = pd.read_excel(
        io.BytesIO(raw),
        sheet_name=None
    )

    result = {}

    for sheet_name, df in sheets.items():

        if df is None:
            continue

        df = clean_dataframe(
            df
        )

        df = convert_dataframe_types(
            df
        )

        if not df.empty:

            result[
                sheet_name
            ] = df

    return result


# =========================================================
# JSON LOADER
# =========================================================

def load_json(
    raw: bytes
):

    try:

        df = pd.read_json(
            io.BytesIO(raw)
        )

    except ValueError:

        import json

        data = json.loads(
            raw.decode(
                "utf-8"
            )
        )

        if isinstance(
            data,
            list
        ):

            df = pd.DataFrame(
                data
            )

        elif isinstance(
            data,
            dict
        ):

            df = pd.DataFrame(
                data
            )

        else:

            raise ValueError(
                "Unsupported JSON structure."
            )

    df = clean_dataframe(
        df
    )

    df = convert_dataframe_types(
        df
    )

    return {
        "data": df
    }


# =========================================================
# PARQUET LOADER
# =========================================================

def load_parquet(
    raw: bytes
):

    df = pd.read_parquet(
        io.BytesIO(raw)
    )

    df = clean_dataframe(
        df
    )

    df = convert_dataframe_types(
        df
    )

    return {
        "data": df
    }


# =========================================================
# PDF LOADER
# =========================================================

def load_pdf(
    raw: bytes
):

    result = {
        "dataframes": {},
        "text": "",
    }

    # -----------------------------------------------------
    # Extract normal PDF text
    # -----------------------------------------------------

    text = extract_pdf_text(
        raw
    )

    result["text"] = text

    # -----------------------------------------------------
    # Extract PDF tables
    # -----------------------------------------------------

    pdf_tables = extract_pdf_tables(
        raw
    )

    valid_table_count = 0

    for (
        page_number,
        table_number,
        df
    ) in pdf_tables:

        if not is_useful_table(
            df
        ):
            continue

        valid_table_count += 1

        key = (
            f"page_{page_number}"
            f"_table_{table_number}"
        )

        result[
            "dataframes"
        ][key] = df

    # -----------------------------------------------------
    # If tables were detected,
    # the DataFrame path will be used
    # by the analytics engine.
    # -----------------------------------------------------

    if valid_table_count > 0:

        result["pdf_tables_detected"] = True
        result["pdf_table_count"] = (
            valid_table_count
        )

    else:

        result["pdf_tables_detected"] = False
        result["pdf_table_count"] = 0

    return result


# =========================================================
# DOCX LOADER
# =========================================================

def load_docx(
    raw: bytes
):

    result = {
        "dataframes": {},
        "text": "",
    }

    document = Document(
        io.BytesIO(raw)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            paragraphs.append(
                text
            )

    result["text"] = "\n".join(
        paragraphs
    )

    # -----------------------------------------------------
    # DOCX TABLES
    # -----------------------------------------------------

    for index, table in enumerate(
        document.tables
    ):

        rows = []

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                row_values.append(
                    cell.text.strip()
                )

            rows.append(
                row_values
            )

        if len(rows) <= 1:
            continue

        header = rows[0]

        # Normalize header
        headers = []

        for column_index, value in enumerate(
            header
        ):

            cleaned = clean_pdf_header(
                value
            )

            if not cleaned:

                cleaned = (
                    f"Column_{column_index + 1}"
                )

            headers.append(
                cleaned
            )

        data_rows = rows[1:]

        df = pd.DataFrame(
            data_rows,
            columns=headers
        )

        df = clean_dataframe(
            df
        )

        df = convert_dataframe_types(
            df
        )

        if not df.empty:

            result[
                "dataframes"
            ][
                f"table_{index + 1}"
            ] = df

    return result


# =========================================================
# TXT LOADER
# =========================================================

def load_txt(
    raw: bytes
):

    return {
        "dataframes": {},
        "text": raw.decode(
            "utf-8",
            errors="ignore"
        ),
    }


# =========================================================
# MAIN FILE LOADER
# =========================================================

def load_file(
    uploaded_file
):

    filename = uploaded_file.name

    extension = get_extension(
        filename
    )

    if extension not in SUPPORTED_EXTENSIONS:

        raise ValueError(
            f"Unsupported file type: "
            f"{extension}"
        )

    raw = uploaded_file.getvalue()

    result = {
        "filename": filename,
        "type": extension,
        "dataframes": {},
        "text": "",
    }

    # =====================================================
    # CSV
    # =====================================================

    if extension == ".csv":

        result[
            "dataframes"
        ] = load_csv(
            raw
        )

    # =====================================================
    # EXCEL
    # =====================================================

    elif extension in [
        ".xlsx",
        ".xls",
    ]:

        result[
            "dataframes"
        ] = load_excel(
            raw
        )

    # =====================================================
    # JSON
    # =====================================================

    elif extension == ".json":

        result[
            "dataframes"
        ] = load_json(
            raw
        )

    # =====================================================
    # PARQUET
    # =====================================================

    elif extension == ".parquet":

        result[
            "dataframes"
        ] = load_parquet(
            raw
        )

    # =====================================================
    # PDF
    # =====================================================

    elif extension == ".pdf":

        pdf_result = load_pdf(
            raw
        )

        result[
            "dataframes"
        ] = pdf_result[
            "dataframes"
        ]

        result[
            "text"
        ] = pdf_result[
            "text"
        ]

        result[
            "pdf_tables_detected"
        ] = pdf_result[
            "pdf_tables_detected"
        ]

        result[
            "pdf_table_count"
        ] = pdf_result[
            "pdf_table_count"
        ]

    # =====================================================
    # DOCX
    # =====================================================

    elif extension == ".docx":

        docx_result = load_docx(
            raw
        )

        result[
            "dataframes"
        ] = docx_result[
            "dataframes"
        ]

        result[
            "text"
        ] = docx_result[
            "text"
        ]

    # =====================================================
    # TXT
    # =====================================================

    elif extension == ".txt":

        txt_result = load_txt(
            raw
        )

        result[
            "text"
        ] = txt_result[
            "text"
        ]

    return result